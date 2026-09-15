"""Passive Candidate Service (Section 4 / Section 3.7).

Scheduled batch job over existing school records. Phase 0 is limited to schools
with grade/attendance data already in DIGITAL (spreadsheet) form — NO OCR
(Section 8). It reads a SchoolRecordImport's referenced file and emits
CandidateSignals.

HARD CONSTRAINTS:
  * Output is a CandidateSignal only — never a DomainScore, never a decision
    (Section 2 / 7). confidence is fixed 'low' by the schema.
  * A signal is a *reason to invite a child to screen*. It cannot advance anyone.
    learner_id is left null here (the Section 5 "stub may not exist yet" case) —
    the pseudonymous school-scoped ref is recorded in evidence_text so the
    invitation can be routed without creating identity prematurely.

This is a functional skeleton: it really parses a spreadsheet and really writes
signals, but the spike heuristics are deliberately simple placeholders to tune
against real pilot data (Section 9 notes this needs a discovery pass).
"""
from __future__ import annotations

import csv
import statistics
import uuid
from collections import defaultdict
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.enums import SignalType
from app.models import CandidateSignal, SchoolRecordImport

# A learner whose subject scores spread at least this wide is worth inviting.
GRADE_SPREAD_THRESHOLD = 30


def _read_grade_rows(path: Path) -> list[dict]:
    """Read (learner_ref, subject, score) rows from a .csv or .xlsx grade sheet."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="") as fh:
            return [dict(r) for r in csv.DictReader(fh)]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:]]
    raise ValueError(f"Unsupported grade sheet format: {path.suffix}")


def run_grade_variance_mining(
    db: Session, import_id: uuid.UUID
) -> list[CandidateSignal]:
    """Emit grade_variance CandidateSignals for one grades import."""
    imp = db.get(SchoolRecordImport, import_id)
    if imp is None:
        raise ValueError(f"No SchoolRecordImport {import_id}")

    path = Path(imp.raw_reference)
    if not path.exists():
        raise FileNotFoundError(
            f"raw_reference file not found: {path} (imports point to files, "
            "never inline data — Section 5)"
        )

    by_learner: dict[str, list[float]] = defaultdict(list)
    for row in _read_grade_rows(path):
        ref = str(row.get("learner_ref", "")).strip()
        try:
            score = float(row.get("score"))
        except (TypeError, ValueError):
            continue
        if ref:
            by_learner[ref].append(score)

    created: list[CandidateSignal] = []
    for ref, scores in by_learner.items():
        if len(scores) < 2:
            continue
        spread = max(scores) - min(scores)
        if spread < GRADE_SPREAD_THRESHOLD:
            continue
        stdev = round(statistics.pstdev(scores), 1)
        signal = CandidateSignal(
            learner_id=None,  # stub not created yet — this is an invitation, not a score
            source_import_id=imp.id,
            signal_type=SignalType.grade_variance,
            evidence_text=(
                f"learner_ref={ref}: subject-score spread {spread:.0f} "
                f"(stdev {stdev}) across {len(scores)} subjects — uneven profile "
                "worth inviting to the active screener. ADVISORY ONLY; not a score."
            ),
            confidence="low",
        )
        db.add(signal)
        created.append(signal)

    db.flush()
    return created
